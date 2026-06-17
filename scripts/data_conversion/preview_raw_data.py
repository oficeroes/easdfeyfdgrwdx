# -*- coding: utf-8 -*-
"""快速预览数据包中的 PDF 说明和 Excel 前几行。

用法：
    python scripts/data_conversion/preview_raw_data.py
    python scripts/data_conversion/preview_raw_data.py --rows 20
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pdfplumber


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "数据包"


def preview_pdf(path: Path, chars: int) -> None:
    print("=" * 70)
    print(f"PDF: {path.name}")
    print("=" * 70)
    try:
        with pdfplumber.open(path) as pdf:
            text = "".join((page.extract_text() or "") + "\n" for page in pdf.pages)
        print((text or "[未提取到文本]").strip()[:chars])
    except Exception as exc:
        print(f"[读取失败] {exc}")
    print()


def preview_excel(path: Path, rows: int) -> None:
    print("=" * 70)
    print(f"Excel: {path.name}")
    print("=" * 70)
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            print(f"  Sheet: {sheet_name} ({worksheet.max_row} 行 x {worksheet.max_column} 列)")
            for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row or 0, rows), values_only=True):
                print(f"    {list(row)}")
            print()
        workbook.close()
    except Exception as exc:
        print(f"[读取失败] {exc}\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="在终端中预览数据包内容。")
    parser.add_argument("--rows", type=int, default=10, help="每个 Excel 工作表预览多少行。")
    parser.add_argument("--chars", type=int, default=3000, help="每个 PDF 预览多少个字符。")
    args = parser.parse_args()

    for pdf_path in sorted(DATA_DIR.glob("*.pdf")):
        preview_pdf(pdf_path, args.chars)
    for xlsx_path in sorted(DATA_DIR.glob("*.xlsx")):
        preview_excel(xlsx_path, args.rows)


if __name__ == "__main__":
    main()
