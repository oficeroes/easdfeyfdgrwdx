# -*- coding: utf-8 -*-
"""将数据包中的 Excel 文件批量转换为 CSV。

用法：
    python scripts/data_conversion/convert_to_csv.py

输出：
    表格数据/*.csv
    表格数据/数据预览/*_预览.csv
    表格数据/_元数据摘要.json
"""

from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "数据包"
OUTPUT_DIR = ROOT / "表格数据"
PREVIEW_DIR = OUTPUT_DIR / "数据预览"

DEFAULT_XLSX_FILES = [
    "2025 團隊賽數據包.xlsx",
    "副本2026 團隊賽數據包.xlsx",
]


def safe_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def write_csv(filepath: Path, headers: list[str], rows: list[list[Any]]) -> None:
    filepath.parent.mkdir(parents=True, exist_ok=True)
    with filepath.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([safe_value(cell) for cell in row])


def convert_workbook(xlsx_name: str, preview_rows: int) -> dict[str, Any] | None:
    input_path = DATA_DIR / xlsx_name
    if not input_path.exists():
        print(f"跳过：未找到 {input_path.relative_to(ROOT)}")
        return None

    print(f"\n处理：{input_path.relative_to(ROOT)}")
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    base_name = input_path.stem
    sheets_meta: list[dict[str, Any]] = []

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        if not rows:
            print(f"  Sheet {sheet_name} 为空，跳过")
            continue

        headers = [safe_value(cell) for cell in rows[0]]
        data_rows = rows[1:]

        csv_filename = f"{base_name}_{sheet_name}.csv"
        preview_filename = f"{base_name}_{sheet_name}_预览.csv"
        write_csv(OUTPUT_DIR / csv_filename, headers, data_rows)
        write_csv(PREVIEW_DIR / preview_filename, headers, data_rows[:preview_rows])

        columns_info = []
        for index, header in enumerate(headers):
            sample_values = [
                row[index]
                for row in data_rows[:200]
                if index < len(row) and row[index] is not None
            ]
            value_types = sorted({type(value).__name__ for value in sample_values})
            null_count = sum(1 for row in data_rows if index >= len(row) or row[index] is None)
            null_pct = round(null_count / len(data_rows) * 100, 1) if data_rows else 0
            columns_info.append(
                {
                    "索引": index,
                    "列名": header,
                    "数据类型": ", ".join(value_types) if value_types else "unknown",
                    "空值比例": null_pct,
                }
            )

        sheets_meta.append(
            {
                "工作表名称": sheet_name,
                "行数_含表头": worksheet.max_row,
                "数据行数": len(data_rows),
                "列数": worksheet.max_column,
                "CSV文件名": csv_filename,
                "预览文件名": preview_filename,
                "列信息": columns_info,
            }
        )
        print(f"  已生成 {csv_filename}，数据行 {len(data_rows)}")

    workbook.close()
    return {
        "源文件": xlsx_name,
        "工作表数量": len(sheets_meta),
        "工作表详情": sheets_meta,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="将数据包中的 xlsx 文件转换为 CSV。")
    parser.add_argument(
        "files",
        nargs="*",
        help="要转换的数据包文件名；不填写时转换 2025 和 2026 默认数据包。",
    )
    parser.add_argument("--preview-rows", type=int, default=100, help="预览 CSV 保留的数据行数。")
    args = parser.parse_args()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)

    target_files = args.files or DEFAULT_XLSX_FILES
    metadata = []
    for xlsx_name in target_files:
        result = convert_workbook(xlsx_name, args.preview_rows)
        if result:
            metadata.append(result)

    metadata_path = OUTPUT_DIR / "_元数据摘要.json"
    with metadata_path.open("w", encoding="utf-8", newline="") as file:
        json.dump(
            {
                "转换时间": datetime.now().isoformat(timespec="seconds"),
                "转换工具": "openpyxl + csv",
                "源目录": str(DATA_DIR.relative_to(ROOT)),
                "输出目录": str(OUTPUT_DIR.relative_to(ROOT)),
                "文件总数": len(target_files),
                "成功": len(metadata),
                "失败": len(target_files) - len(metadata),
                "文件详情": metadata,
            },
            file,
            ensure_ascii=False,
            indent=2,
        )

    print(f"\n转换完成，元数据写入：{metadata_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
