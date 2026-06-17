# -*- coding: utf-8 -*-
"""将数据包中的 PDF 和 Excel 导出为便于阅读的文本。

用法：
    python scripts/data_conversion/export_readable_text.py

输出：
    可读文本/*.txt
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pdfplumber


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "数据包"
OUTPUT_DIR = ROOT / "可读文本"


def write_text_file(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8-sig")


def export_pdf_text(pdf_path: Path) -> None:
    output_path = OUTPUT_DIR / f"{pdf_path.stem}.txt"
    try:
        with pdfplumber.open(pdf_path) as pdf:
            text = "".join((page.extract_text() or "") + "\n" for page in pdf.pages)
        write_text_file(output_path, text if text.strip() else "[此 PDF 未提取到文本内容]")
        print(f"已导出 PDF 文本：{output_path.relative_to(ROOT)}")
    except Exception as exc:
        write_text_file(output_path, f"[读取失败] {exc}")
        print(f"PDF 读取失败：{pdf_path.name} -> {exc}")


def export_excel_text(xlsx_path: Path, max_rows: int | None) -> None:
    try:
        workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            lines = [
                f"文件: {xlsx_path.name}",
                f"工作表: {sheet_name}",
                f"行数: {worksheet.max_row}",
                f"列数: {worksheet.max_column}",
                "",
            ]
            for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if max_rows and index > max_rows:
                    lines.append(f"... 已截断，仅导出前 {max_rows} 行")
                    break
                lines.append("\t".join("" if value is None else str(value) for value in row))

            output_path = OUTPUT_DIR / f"{xlsx_path.stem}_{sheet_name}.txt"
            write_text_file(output_path, "\n".join(lines))
            print(f"已导出 Excel 文本：{output_path.relative_to(ROOT)}")
        workbook.close()
    except Exception as exc:
        output_path = OUTPUT_DIR / f"{xlsx_path.stem}_ERROR.txt"
        write_text_file(output_path, f"[读取失败] {exc}")
        print(f"Excel 读取失败：{xlsx_path.name} -> {exc}")


def main() -> None:
    parser = argparse.ArgumentParser(description="把数据包中的 PDF / Excel 导出为文本文件。")
    parser.add_argument("--max-rows", type=int, default=0, help="每个 Excel 工作表最多导出多少行；0 表示不限制。")
    args = parser.parse_args()

    pdf_files = sorted(DATA_DIR.glob("*.pdf"))
    xlsx_files = sorted(DATA_DIR.glob("*.xlsx"))
    if not pdf_files and not xlsx_files:
        raise SystemExit(f"未在 {DATA_DIR} 找到 PDF 或 Excel 数据包。")

    for pdf_path in pdf_files:
        export_pdf_text(pdf_path)
    for xlsx_path in xlsx_files:
        export_excel_text(xlsx_path, args.max_rows or None)

    print(f"\n全部导出完成，输出目录：{OUTPUT_DIR.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
