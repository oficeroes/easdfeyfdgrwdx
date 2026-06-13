"""
============================================================
澳门生物多样性团队赛 — xlsx → CSV 批量转换工具
============================================================
将数据包中的所有 .xlsx 文件转换为 AI 可读的：
  1. CSV 文件（每个 Sheet 一个）
  2. JSON 元数据摘要文件
  3. 前 100 行预览文件（方便 AI 快速理解数据结构）

输出目录：项目根目录下的 表格数据/
============================================================
"""

import csv
import json
import os
import sys
from datetime import datetime

import openpyxl


# ============================================================
# 路径配置
# ============================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, '数据包')
OUTPUT_DIR = os.path.join(BASE_DIR, '表格数据')
PREVIEW_DIR = os.path.join(OUTPUT_DIR, '数据预览')

# 要处理的 xlsx 文件列表
XLSX_FILES = [
    '2025 團隊賽數據包.xlsx',
    '副本2026 團隊賽數據包.xlsx',
]


# ============================================================
# 工具函数
# ============================================================
def ensure_dir(path: str):
    """创建目录（如不存在）"""
    os.makedirs(path, exist_ok=True)


def safe_value(val) -> str:
    """将单元格值转为字符串，处理 None 和日期时间"""
    if val is None:
        return ''
    if isinstance(val, datetime):
        return val.isoformat()
    return str(val)


def csv_quote_if_needed(val: str) -> str:
    """如果值含逗号/换行/引号，用 CSV 标准引号包裹"""
    if ',' in val or '\n' in val or '"' in val:
        return '"' + val.replace('"', '""') + '"'
    return val


def write_csv(filepath: str, headers: list[str], rows: list[list]):
    """写入 CSV 文件（UTF-8 BOM，Excel 兼容）"""
    with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([safe_value(cell) for cell in row])


# ============================================================
# 核心逻辑
# ============================================================
def convert_xlsx_to_csv(xlsx_name: str) -> dict | None:
    """
    转换单个 xlsx 文件为 CSV。
    返回元数据字典，失败返回 None。
    """
    input_path = os.path.join(DATA_DIR, xlsx_name)

    if not os.path.exists(input_path):
        print(f'  ⚠ 文件不存在，跳过: {xlsx_name}')
        return None

    print(f'\n📄 正在处理: {xlsx_name}')
    try:
        wb = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f'  ❌ 无法打开文件: {e}')
        return None

    base_name = os.path.splitext(xlsx_name)[0]
    sheets_meta = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        if max_row == 0:
            print(f'  📋 Sheet "{sheet_name}" 为空，跳过')
            continue

        print(f'  📋 Sheet "{sheet_name}" → {max_row} 行 × {max_col} 列，正在读取...')

        # 读取所有数据
        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))

        if not all_rows:
            continue

        headers = [safe_value(h) for h in all_rows[0]]
        data_rows = all_rows[1:]

        # —— 写入完整 CSV ——
        csv_filename = f'{base_name}_{sheet_name}.csv'
        csv_path = os.path.join(OUTPUT_DIR, csv_filename)
        write_csv(csv_path, headers, data_rows)
        file_size_mb = os.path.getsize(csv_path) / (1024 * 1024)
        print(f'     ✅ 完整 CSV: {csv_filename} ({file_size_mb:.2f} MB)')

        # —— 写入预览 CSV（前 100 行数据）——
        preview_filename = f'{base_name}_{sheet_name}_预览.csv'
        preview_path = os.path.join(PREVIEW_DIR, preview_filename)
        preview_rows = data_rows[:100]
        write_csv(preview_path, headers, preview_rows)
        print(f'     📝 预览 CSV: {preview_filename} ({len(preview_rows)} 行)')

        # 收集列信息
        columns_info = []
        for idx, h in enumerate(headers):
            # 采样非空值来推断数据类型
            sample_values = []
            for row in data_rows[:200]:
                val = row[idx] if idx < len(row) else None
                if val is not None:
                    sample_values.append(val)
            
            col_type = 'unknown'
            if sample_values:
                types = set(type(v).__name__ for v in sample_values)
                col_type = ', '.join(sorted(types))
            
            # 统计空值比例
            total = len(data_rows)
            null_count = sum(1 for row in data_rows if idx >= len(row) or row[idx] is None)
            null_pct = round(null_count / total * 100, 1) if total > 0 else 0

            columns_info.append({
                '索引': idx,
                '列名': h,
                '数据类型': col_type,
                '空值比例%': null_pct,
            })

        sheets_meta.append({
            '工作表名称': sheet_name,
            '行数（含表头）': max_row,
            '数据行数': len(data_rows),
            '列数': max_col,
            'CSV文件名': csv_filename,
            '预览文件名': preview_filename,
            '列信息': columns_info,
        })

    wb.close()

    return {
        '源文件': xlsx_name,
        '工作表数量': len(sheets_meta),
        '工作表详情': sheets_meta,
    }


def main():
    print('=' * 60)
    print('🐍 澳门生物多样性 — xlsx → CSV 批量转换')
    print('=' * 60)

    ensure_dir(OUTPUT_DIR)
    ensure_dir(PREVIEW_DIR)

    all_metadata = []
    success_count = 0
    fail_count = 0

    for xlsx_name in XLSX_FILES:
        meta = convert_xlsx_to_csv(xlsx_name)
        if meta:
            all_metadata.append(meta)
            success_count += 1
        else:
            fail_count += 1

    # —— 写入 JSON 元数据 ——
    metadata_path = os.path.join(OUTPUT_DIR, '_元数据摘要.json')
    with open(metadata_path, 'w', encoding='utf-8-sig') as f:
        json.dump({
            '转换时间': datetime.now().isoformat(),
            '转换工具': 'openpyxl + csv (Python)',
            '文件总数': len(XLSX_FILES),
            '成功': success_count,
            '失败': fail_count,
            '输出目录': OUTPUT_DIR,
            '文件详情': all_metadata,
        }, f, ensure_ascii=False, indent=2)
    print(f'\n📋 元数据摘要: {metadata_path}')

    # —— 终局报告 ——
    print('\n' + '=' * 60)
    print('✅ 转换完成！')
    print(f'   成功: {success_count} 个文件')
    print(f'   失败: {fail_count} 个文件')
    print(f'   输出目录: {OUTPUT_DIR}')
    print(f'   预览目录: {PREVIEW_DIR}')
    print('=' * 60)

    # 列出输出文件
    print('\n📁 输出文件清单:')
    for root, dirs, files in os.walk(OUTPUT_DIR):
        for fname in sorted(files):
            fpath = os.path.join(root, fname)
            size_kb = os.path.getsize(fpath) / 1024
            rel = os.path.relpath(fpath, BASE_DIR)
            print(f'   {size_kb:8.1f} KB  {rel}')


if __name__ == '__main__':
    os.chdir(BASE_DIR)
    main()
