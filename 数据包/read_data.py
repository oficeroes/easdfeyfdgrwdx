import openpyxl, pdfplumber, os

os.chdir(os.path.dirname(os.path.abspath(__file__)))
results = {}

# --- 读取 PDF 说明文件 ---
for pdf_name in ['2025 團隊賽數據包說明.pdf', '2026 團隊賽數據包說明.pdf']:
    try:
        with pdfplumber.open(pdf_name) as pdf:
            text = ''
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + '\n'
            results[pdf_name] = text[:3000]
    except Exception as e:
        results[pdf_name] = f'[读取失败] {e}'

# --- 读取 Excel 数据文件 ---
for xls_name in ['2025 團隊賽數據包.xlsx', '2026 團隊賽數據包.xlsx', '副本2026 團隊賽數據包.xlsx']:
    try:
        wb = openpyxl.load_workbook(xls_name, data_only=True)
        info = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 30), values_only=True):
                rows.append(list(r))
            info[sheet_name] = {
                'rows_count': ws.max_row,
                'cols_count': ws.max_column,
                'preview': rows[:10]
            }
        results[xls_name] = info
    except Exception as e:
        results[xls_name] = f'[读取失败] {e}'

# 输出
for name, data in results.items():
    print('=' * 60)
    print(f'### {name}')
    print('=' * 60)
    if isinstance(data, str):
        print(data)
    elif isinstance(data, dict):
        for sname, sinfo in data.items():
            print(f'  Sheet: {sname} ({sinfo["rows_count"]}行 x {sinfo["cols_count"]}列)')
            for row in sinfo['preview']:
                print(f'    {row}')
    print()
