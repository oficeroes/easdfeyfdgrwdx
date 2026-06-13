import os
import openpyxl
import pdfplumber


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, '数据包')
OUTPUT_DIR = os.path.join(BASE_DIR, '可读文本')


def ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def write_text_file(path, content):
    with open(path, 'w', encoding='utf-8-sig') as f:
        f.write(content)


def export_pdf_text(pdf_name):
    input_path = os.path.join(DATA_DIR, pdf_name)
    output_path = os.path.join(OUTPUT_DIR, f'{os.path.splitext(pdf_name)[0]}.txt')
    try:
        with pdfplumber.open(input_path) as pdf:
            text = ''.join((page.extract_text() or '') + '\n' for page in pdf.pages)
        if not text.strip():
            text = '[此 PDF 没有提取到文本内容]'
        write_text_file(output_path, text)
        print(f'已导出 PDF 文本: {output_path}')
    except Exception as e:
        write_text_file(output_path, f'[读取失败] {e}')
        print(f'PDF 读取失败: {pdf_name} -> {e}')


def export_excel_text(xls_name):
    input_path = os.path.join(DATA_DIR, xls_name)
    try:
        wb = openpyxl.load_workbook(input_path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines = [
                f'文件: {xls_name}',
                f'工作表: {sheet_name}',
                f'行数: {ws.max_row}',
                f'列数: {ws.max_column}',
                '',
            ]
            for row in ws.iter_rows(values_only=True):
                lines.append('\t'.join('' if value is None else str(value) for value in row))

            out_name = f'{os.path.splitext(xls_name)[0]}_{sheet_name}.txt'
            out_path = os.path.join(OUTPUT_DIR, out_name)
            write_text_file(out_path, '\n'.join(lines))
            print(f'已导出 Excel 文本: {out_path}')
    except Exception as e:
        error_path = os.path.join(OUTPUT_DIR, f'{os.path.splitext(xls_name)[0]}_ERROR.txt')
        write_text_file(error_path, f'[读取失败] {e}')
        print(f'Excel 读取失败: {xls_name} -> {e}')


def main():
    ensure_output_dir()
    os.chdir(BASE_DIR)

    pdf_files = ['2025 團隊賽數據包說明.pdf', '2026 團隊賽數據包說明.pdf']
    xls_files = ['2025 團隊賽數據包.xlsx', '2026 團隊賽數據包.xlsx', '副本2026 團隊賽數據包.xlsx']

    for pdf_name in pdf_files:
        export_pdf_text(pdf_name)

    for xls_name in xls_files:
        export_excel_text(xls_name)

    print('\n全部转换完成。文本文件保存在:')
    print(OUTPUT_DIR)


if __name__ == '__main__':
    main()
